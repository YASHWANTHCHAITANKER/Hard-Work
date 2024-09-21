from docx import Document
from openpyxl import Workbook

def convert_docx_to_excel(docx_file, excel_file):
    # Load the Word document
    doc = Document(docx_file)
    
    # Create a new Excel workbook
    wb = Workbook()
    ws = wb.active
    
    # Iterate through paragraphs in the Word document
    for paragraph in doc.paragraphs:
        # Split paragraph text by some delimiter, e.g., space
        data = paragraph.text.split()
        # Write the data to the Excel sheet row-wise
        ws.append(data)
    
    # Save the Excel workbook
    wb.save(excel_file)

# Example usage
docx_file = "input.docx"
excel_file = "example.xlsx"
convert_docx_to_excel(docx_file, excel_file)
