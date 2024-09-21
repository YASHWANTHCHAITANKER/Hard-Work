from docx import Document
from docx.shared import Inches
import pandas as pd

def create_word_document(data, word_file):
    # Create a new Word document
    doc = Document()

    # Add a table to the Word document
    rows = len(data)
    cols = len(data[0])
    table = doc.add_table(rows=rows, cols=cols)

    # Add data to the table
    for i in range(rows):
        for j in range(cols):
            table.cell(i, j).text = str(data[i][j])

    # Save the Word document
    doc.save(word_file)
    print(f"Word document saved as '{word_file}'.")

# Read input data from Excel file
def read_excel_data(excel_file):
    df = pd.read_excel(excel_file)
    data = df.values.tolist()
    return data

# Example usage
input_excel_file = "Book1.xlsx"
output_word_file = "output.docx"

input_data = read_excel_data(input_excel_file)
create_word_document(input_data, output_word_file)
